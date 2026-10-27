#!/usr/bin/env python3
"""매달 받는 전체 신청명단 raw data를 정리된 양식으로 변환한다.

사용법:
    python3 scripts/process_apply_list.py <신청명단.xlsx> <환급과정_커리큘럼.xlsx> <비환급과정_커리큘럼.xlsx> <출력파일.xlsx>
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pricing import (  # noqa: E402
    added_note, classify_course, clean_company_name, filter_excluded_rows, load_price_table, read_apply_list,
    split_tel_rows, style_worksheet,
)

REQUIRED_COLUMNS = [
    "순번", "신청ID", "구분", "출처", "회사명", "지사구분", "성함", "아이디",
    "생년월일", "휴대전화", "메일주소", "산인공코드", "신청과정명",
    "교육시작일", "교육종료일", "신청일시", "결제수단", "결제일", "결제금액", "결제상태", "신청상태",
]

# 정리된 신청명단에는 필요 없어서 결과 파일에서 빼는 칸
DROP_FROM_OUTPUT = ["신청ID", "출처", "회사명", "지사구분", "결제수단", "결제일", "결제금액", "결제상태"]

WHITESPACE_CHECK_COLUMNS = ["성함", "아이디", "휴대전화", "메일주소"]


def has_whitespace_issue(value) -> bool:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return False
    text = str(value)
    if text != text.strip():
        return True
    if "  " in text:
        return True
    return False


def run(apply_list_path: str, refund_curriculum_path: str, general_price_path: str, output_path: str,
        added_path: str = None) -> dict:
    """신청명단을 정리해 output_path에 저장하고 요약 결과를 돌려준다.
    added_path(추가 신청 파일)를 주면 원본 + 추가 신청을 합쳐서 정리한다.
    입력 파일에 필요한 칸이 없으면 ValueError를 발생시킨다(추측해서 넘어가지 않음)."""
    df = read_apply_list(apply_list_path, REQUIRED_COLUMNS, added_path)

    df, excluded_df, excluded_count = filter_excluded_rows(df)
    df, tel_df = split_tel_rows(df)

    refund_table = load_price_table(refund_curriculum_path)
    general_table = load_price_table(general_price_path)

    df["기업명"] = df["회사명"].apply(clean_company_name)

    price_results = df.apply(
        lambda r: classify_course(r["신청과정명"], refund_table, general_table, r["구분"]), axis=1
    )
    df["환급구분"] = price_results.apply(lambda r: r["환급구분"])
    df["수강료"] = price_results.apply(lambda r: r["수강료"])
    df["정부지원금"] = price_results.apply(lambda r: r["정부지원금"])
    df["기업부담금"] = price_results.apply(lambda r: r["기업부담금"])
    price_notes = price_results.apply(lambda r: r["확인필요"])

    def whitespace_note(row) -> str:
        issues = [col for col in WHITESPACE_CHECK_COLUMNS if has_whitespace_issue(row.get(col))]
        return (", ".join(issues) + " 공백 확인필요") if issues else ""

    whitespace_notes = df.apply(whitespace_note, axis=1)
    added_notes = df["_추가"].apply(added_note)
    df["확인필요"] = [
        " / ".join(n for n in (ws_note, price_note, dup_note, add_note) if n)
        for ws_note, price_note, dup_note, add_note in zip(whitespace_notes, price_notes, df["_중복메모"], added_notes)
    ]

    base_columns = [c for c in REQUIRED_COLUMNS if c not in DROP_FROM_OUTPUT]
    insert_at = base_columns.index("구분") + 1
    ordered_columns = (base_columns[:insert_at] + ["기업명"] + base_columns[insert_at:]
                        + ["환급구분", "수강료", "정부지원금", "기업부담금", "확인필요"])
    added_count = int(df["_추가"].sum())
    duplicate_count = int((df["_중복메모"] != "").sum())
    df = df[ordered_columns]

    df.to_excel(output_path, index=False, sheet_name="정리된_신청명단")

    wb = load_workbook(output_path)
    ws = wb["정리된_신청명단"]
    fill = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
    check_col_idx = ordered_columns.index("확인필요") + 1
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=check_col_idx)
        if cell.value:
            cell.fill = fill
    style_worksheet(ws, currency_columns=["수강료", "정부지원금", "기업부담금"], text_columns=["아이디"])
    wb.save(output_path)

    # 5번 파일: 프리패스/테스트/카드결제 때문에 결과에서 빠진 신청 건만 모아서 따로 저장.
    output_path = Path(output_path)
    name_suffix = output_path.stem
    if name_suffix.startswith("1_정리된_신청명단"):
        name_suffix = name_suffix[len("1_정리된_신청명단"):]
    excluded_output_path = output_path.parent / f"5_제외된_신청명단{name_suffix}.xlsx"
    excluded_columns = [c for c in REQUIRED_COLUMNS if c in excluded_df.columns] + ["제외사유"]
    excluded_df[excluded_columns].to_excel(excluded_output_path, index=False, sheet_name="제외된_신청명단")

    excluded_wb = load_workbook(excluded_output_path)
    style_worksheet(excluded_wb["제외된_신청명단"], text_columns=["아이디"])
    excluded_wb.save(excluded_output_path)

    total = len(df)
    refund_count = int((df["환급구분"] == "환급").sum())
    nonrefund_count = int((df["환급구분"] == "비환급").sum())
    flagged = int((df["확인필요"] != "").sum())

    return {
        "total": total,
        "refund": refund_count,
        "nonrefund": nonrefund_count,
        "flagged": flagged,
        "excluded": excluded_count,
        "excluded_output_path": str(excluded_output_path),
        "tel_excluded": len(tel_df),
        "added": added_count,
        "duplicates": duplicate_count,
        "output_path": str(output_path),
    }


def main():
    if len(sys.argv) != 5:
        print("사용법: python3 scripts/process_apply_list.py "
              "<신청명단.xlsx> <환급과정_커리큘럼.xlsx> <비환급과정_커리큘럼.xlsx> <출력파일.xlsx>")
        sys.exit(1)

    apply_list_path, refund_curriculum_path, general_price_path, output_path = sys.argv[1:5]
    try:
        result = run(apply_list_path, refund_curriculum_path, general_price_path, output_path)
    except ValueError as e:
        print(str(e))
        sys.exit(1)

    print(f"완료: 총 {result['total']}건 (프리패스/테스트/카드결제 {result['excluded']}건, "
          f"도쿄일렉트론코리아 {result['tel_excluded']}건 제외) / "
          f"환급 {result['refund']}건 / 비환급 {result['nonrefund']}건 / 확인필요 {result['flagged']}건")
    print(f"결과 파일: {result['output_path']}")
    print(f"제외된 신청명단 파일: {result['excluded_output_path']}")


if __name__ == "__main__":
    main()
