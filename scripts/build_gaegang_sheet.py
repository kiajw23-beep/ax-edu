#!/usr/bin/env python3
"""매달 받는 전체 신청명단을 "개강내역시트"에 채울 내용(기업·과정별 집계)으로 만든다.

사용법:
    python3 scripts/build_gaegang_sheet.py <신청명단.xlsx> <환급과정_커리큘럼.xlsx> <비환급과정_커리큘럼.xlsx> <출력파일.xlsx>

출력 칸(사용자가 지정한 순서 그대로):
    기업명, 개강일, 교육종료일, 교육회차, 산인공코드, 과정명, 인원,
    단가 (1인기준), 정부지원금, 기업부담금, 총 금액

규칙:
- 기업명 + 과정명 + 개강일 + 교육종료일이 같은 신청 건을 하나의 줄로 묶는다
  (인원 = 그 줄에 묶인 학생 수).
- 단가(1인기준)/정부지원금/기업부담금 계산 규칙은 generate_company_files.py와 동일
  (환급·비환급과정 커리큘럼 대조, 양쪽에 다 있으면 환급 우선 + 확인필요 메모).
  정부지원금/기업부담금/총 금액은 인원을 곱한 "그룹 합계" 금액이다.
- 교육회차 칸은 이 도구가 알 수 없는 정보라 비워둔다(사용자가 직접 채움). 추가 신청 반영 시
  기존 개강내역 파일에 채워둔 교육회차는 그대로 옮긴다.
- 산인공코드가 같은 그룹 안에서 서로 다르면 비워두고 확인필요로 표시한다.
- 한 그룹 안에 금액을 알 수 없는(확인필요) 학생이 섞여 있으면, 그 그룹의 금액
  전체를 비워두고 확인필요로 표시한다(일부만 추정해서 채우지 않는다).
"""
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pricing import (  # noqa: E402
    classify_course, clean_company_name, filter_excluded_rows, load_price_table, normalize, read_apply_list,
    split_tel_rows, style_worksheet,
)

REQUIRED_APPLY_COLUMNS = [
    "순번", "신청ID", "구분", "출처", "회사명", "지사구분", "성함", "아이디",
    "생년월일", "휴대전화", "메일주소", "산인공코드", "신청과정명",
    "교육시작일", "교육종료일", "신청일시", "결제수단", "결제일", "결제금액", "결제상태", "신청상태",
]

OUTPUT_COLUMNS = [
    "기업명", "개강일", "교육종료일", "교육회차", "산인공코드", "과정명", "인원",
    "단가 (1인기준)", "정부지원금", "기업부담금", "총 금액", "확인필요",
]


def sheet_key(company, course, start_date, end_date) -> tuple:
    return tuple(normalize(v) for v in (company, course, start_date, end_date))


def load_existing_sheet(path: str) -> dict:
    """기존 개강내역 파일에서 (기업명, 과정명, 개강일, 교육종료일) -> {교육회차, 인원}을 읽는다."""
    old_df = pd.read_excel(path, sheet_name="개강내역")
    missing = [c for c in OUTPUT_COLUMNS if c not in old_df.columns]
    if missing:
        raise ValueError(f"기존 개강내역 파일({path})에 다음 칸이 없습니다: {missing}")
    table = {}
    for _, r in old_df.iterrows():
        session = r["교육회차"]
        table[sheet_key(r["기업명"], r["과정명"], r["개강일"], r["교육종료일"])] = {
            "교육회차": "" if pd.isna(session) else session,
            "인원": r["인원"],
        }
    return table


def run(apply_path: str, refund_curriculum_path: str, general_price_path: str, output_path: str,
        added_path: str = None, existing_path: str = None) -> dict:
    """개강내역 집계 결과를 output_path에 저장하고 요약 결과를 돌려준다.
    added_path(추가 신청 파일)를 주면 원본 + 추가 신청을 합쳐서 다시 집계한다.
    existing_path(기존 개강내역 파일)를 주면 직접 채운 교육회차를 그대로 옮기고,
    인원이 바뀐 줄·새로 생긴 줄은 확인필요(노란색)에 표시한다.
    입력 파일에 필요한 칸이 없으면 ValueError를 발생시킨다(추측해서 넘어가지 않음)."""
    df = read_apply_list(apply_path, REQUIRED_APPLY_COLUMNS, added_path)

    df, _, excluded_count = filter_excluded_rows(df)
    df, tel_df = split_tel_rows(df)

    refund_table = load_price_table(refund_curriculum_path)
    general_table = load_price_table(general_price_path)

    per_student = []
    for _, r in df.iterrows():
        result = classify_course(r["신청과정명"], refund_table, general_table)
        per_student.append({
            "기업명": clean_company_name(r["회사명"]),
            "개강일": r["교육시작일"],
            "교육종료일": r["교육종료일"],
            "과정명": r["신청과정명"],
            "산인공코드": r["산인공코드"],
            "단가": result["수강료"],
            "1인정부지원금": result["정부지원금"],
            "1인기업부담금": result["기업부담금"],
            "확인필요": " / ".join(n for n in (result["확인필요"], r["_중복메모"]) if n),
        })
    per_student_df = pd.DataFrame(per_student)

    group_keys = ["기업명", "과정명", "개강일", "교육종료일"]
    existing = load_existing_sheet(existing_path) if existing_path else None
    out_rows = []
    for key, group in per_student_df.groupby(group_keys, sort=False, dropna=False):
        company, course, start_date, end_date = key
        headcount = len(group)
        notes = []

        unit_prices = group["단가"].dropna().unique()
        price_known = group["단가"].notna().all() and len(unit_prices) == 1
        for student_note in group["확인필요"]:
            for note in student_note.split(" / ") if student_note else []:
                if note not in notes:
                    notes.append(note)

        codes = group["산인공코드"].dropna().unique()
        if len(codes) > 1:
            code = None
            notes.append(f"산인공코드가 서로 달라({list(codes)}) 확인 필요")
        else:
            code = codes[0] if len(codes) == 1 else None

        if price_known:
            unit_price = unit_prices[0]
            gov_total = group["1인정부지원금"].iloc[0] * headcount
            company_total = group["1인기업부담금"].iloc[0] * headcount
            total_amount = gov_total + company_total
        else:
            unit_price = gov_total = company_total = total_amount = None
            if not any("과정명이 없음" in n or "중복" in n for n in notes):
                notes.append("같은 과정 안에 금액이 서로 다른 학생이 섞여 있어 금액 확인 필요")

        session = ""
        if existing is not None:
            old = existing.get(sheet_key(company, course, start_date, end_date))
            if old is None:
                notes.append("추가 신청으로 새로 생긴 줄")
            else:
                session = old["교육회차"]
                if old["인원"] != headcount:
                    notes.append(f"추가 신청 반영(인원 {old['인원']}→{headcount})")

        out_rows.append({
            "기업명": company,
            "개강일": start_date,
            "교육종료일": end_date,
            "교육회차": session,
            "산인공코드": code,
            "과정명": course,
            "인원": headcount,
            "단가 (1인기준)": unit_price,
            "정부지원금": gov_total,
            "기업부담금": company_total,
            "총 금액": total_amount,
            "확인필요": " / ".join(notes),
        })

    out_df = pd.DataFrame(out_rows, columns=OUTPUT_COLUMNS)
    out_df.to_excel(output_path, index=False, sheet_name="개강내역")

    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    wb = load_workbook(output_path)
    ws = wb["개강내역"]
    fill = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
    check_col_idx = OUTPUT_COLUMNS.index("확인필요") + 1
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=check_col_idx)
        if cell.value:
            cell.fill = fill
    style_worksheet(ws, currency_columns=["단가 (1인기준)", "정부지원금", "기업부담금", "총 금액"])
    wb.save(output_path)

    total_rows = len(out_df)
    flagged = int((out_df["확인필요"] != "").sum())

    return {
        "rows": total_rows,
        "students": len(per_student_df),
        "flagged": flagged,
        "excluded": excluded_count,
        "tel_excluded": len(tel_df),
        "output_path": output_path,
    }


def main():
    if len(sys.argv) != 5:
        print("사용법: python3 scripts/build_gaegang_sheet.py "
              "<신청명단.xlsx> <환급과정_커리큘럼.xlsx> <비환급과정_커리큘럼.xlsx> <출력파일.xlsx>")
        sys.exit(1)

    apply_path, refund_curriculum_path, general_price_path, output_path = sys.argv[1:5]
    try:
        result = run(apply_path, refund_curriculum_path, general_price_path, output_path)
    except ValueError as e:
        print(str(e))
        sys.exit(1)

    print(f"완료: 개강내역 {result['rows']}줄 (신청 {result['students']}건을 묶음, "
          f"프리패스/테스트/카드결제 {result['excluded']}건, "
          f"도쿄일렉트론코리아 {result['tel_excluded']}건 제외) / 확인필요 {result['flagged']}줄")
    print(f"결과 파일: {result['output_path']}")


if __name__ == "__main__":
    main()
