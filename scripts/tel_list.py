#!/usr/bin/env python3
"""신청명단에서 도쿄일렉트론코리아 소속 신청 건만 모아 별도 파일로 만든다.

사용법:
    python3 scripts/tel_list.py <신청명단.xlsx> <환급과정_커리큘럼.xlsx> <비환급_가격표.xlsx> <학점파일.xlsx> <출력파일.xlsx>

출력 칸(사용자가 지정한 순서 그대로):
    성함, 아이디, 과정명, 신청일, 교육시작일, 교육종료일, 학점, 교육비, 확인필요

규칙:
- 프리패스/테스트/카드결제 제외 규칙을 먼저 적용한 뒤, 회사명이
  도쿄일렉트론코리아(법인형태·띄어쓰기 무시)인 건만 남긴다.
- 과정명은 신청과정명 원본 그대로, 신청일은 신청일시 원본 값 그대로 쓴다.
- 학점: 신청과정명을 학점 파일의 과정명 칸들("과정명", "최종 과정명", 제목 없는 칸)과
  공백 차이만 보정해 그대로 대조한다. 못 찾거나 같은 이름이 서로 다른 학점으로
  중복돼 있으면 비워두고 확인필요로 표시한다.
- 교육비: 1인 수강료(1·2·3단계와 같은 환급/비환급 가격표 규칙).
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pricing import classify_course, filter_excluded_rows, find_column, load_price_table, normalize, split_tel_rows  # noqa: E402
from process_apply_list import REQUIRED_COLUMNS  # noqa: E402

OUTPUT_COLUMNS = ["성함", "아이디", "과정명", "신청일", "교육시작일", "교육종료일", "학점", "교육비", "확인필요"]


def load_credit_table(path: str) -> dict:
    """과정명 -> 학점 사전. 같은 이름이 서로 다른 학점으로 있으면 값을 None으로 둔다."""
    df = pd.read_csv(path) if path.lower().endswith(".csv") else pd.read_excel(path)

    credit_col = find_column(df.columns, "학점")
    name_cols = [c for c in df.columns
                 if isinstance(c, str) and ("과정명" in c or c.startswith("Unnamed"))]
    if credit_col is None or not name_cols:
        raise ValueError(f"학점 파일({path})에서 '과정명' 칸이나 '학점' 칸을 찾을 수 없습니다.")

    table = {}
    for _, row in df.iterrows():
        credit = row[credit_col]
        if pd.isna(credit):
            continue
        for col in name_cols:
            if pd.isna(row[col]):
                continue
            name = normalize(row[col])
            if name in table and table[name] != credit:
                table[name] = None
            elif name not in table:
                table[name] = credit
    return table


def run(apply_path: str, refund_curriculum_path: str, general_price_path: str,
        credit_path: str, output_path: str) -> dict:
    df = pd.read_excel(apply_path)
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"신청명단 파일에 다음 칸이 없습니다: {missing}")

    df, _ = filter_excluded_rows(df)
    _, tel_df = split_tel_rows(df)

    refund_table = load_price_table(refund_curriculum_path)
    general_table = load_price_table(general_price_path)
    credit_table = load_credit_table(credit_path)

    rows = []
    for _, r in tel_df.iterrows():
        course = r["신청과정명"]
        price = classify_course(course, refund_table, general_table)
        notes = [price["확인필요"]] if price["확인필요"] else []

        course_norm = normalize(course)
        credit = credit_table.get(course_norm)
        if course_norm not in credit_table:
            notes.append("학점 파일에 과정명이 없음 - 학점 확인 필요")
        elif credit is None:
            notes.append("학점 파일에 같은 과정명이 서로 다른 학점으로 있음 - 학점 확인 필요")

        rows.append({
            "성함": r["성함"],
            "아이디": r["아이디"],
            "과정명": course,
            "신청일": r["신청일시"],
            "교육시작일": r["교육시작일"],
            "교육종료일": r["교육종료일"],
            "학점": credit,
            "교육비": price["수강료"],
            "확인필요": " / ".join(notes),
        })

    out_df = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
    out_df.to_excel(output_path, index=False, sheet_name="도쿄일렉트론코리아")

    wb = load_workbook(output_path)
    ws = wb["도쿄일렉트론코리아"]
    fill = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
    check_col_idx = OUTPUT_COLUMNS.index("확인필요") + 1
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=check_col_idx)
        if cell.value:
            cell.fill = fill
    wb.save(output_path)

    return {
        "total": len(out_df),
        "flagged": int((out_df["확인필요"] != "").sum()),
        "output_path": output_path,
    }


def main():
    if len(sys.argv) != 6:
        print("사용법: python3 scripts/tel_list.py "
              "<신청명단.xlsx> <환급과정_커리큘럼.xlsx> <비환급_가격표.xlsx> <학점파일.xlsx> <출력파일.xlsx>")
        sys.exit(1)

    try:
        result = run(*sys.argv[1:6])
    except ValueError as e:
        print(str(e))
        sys.exit(1)

    print(f"완료: 도쿄일렉트론코리아 {result['total']}건 / 확인필요 {result['flagged']}건")
    print(f"결과 파일: {result['output_path']}")


if __name__ == "__main__":
    main()
